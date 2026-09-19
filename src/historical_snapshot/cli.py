from __future__ import annotations

import argparse
import csv
import json
import sys
from datetime import date, timedelta
from decimal import Decimal
from pathlib import Path

from historical_snapshot.core.bands import DEFAULT_BANDS
from historical_snapshot.core.snapshot import multi_snapshot_to_text
from historical_snapshot.service import DEFAULT_DATA_ROOT, run_snapshot, sync_properties


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        prog="snapshot",
        description="Get property-level historical performance snapshots from CSV or Google Sheets data.",
    )
    subparsers = parser.add_subparsers(dest="command")

    sync_parser = subparsers.add_parser(
        "sync",
        help="Pull property tabs from Google Sheets into the local cache",
    )
    sync_parser.add_argument(
        "--property-folder",
        default=None,
        help="Sync one property by config folder name (e.g. onera, lafave)",
    )
    sync_parser.add_argument(
        "--data-root",
        default=str(DEFAULT_DATA_ROOT),
        help="Data root directory (default: data)",
    )
    sync_parser.add_argument(
        "--format",
        choices=("text", "json"),
        default="text",
        help="Output format",
    )

    run_parser = subparsers.add_parser("run", help="Run a snapshot")
    _add_run_arguments(run_parser)

    export_parser = subparsers.add_parser("export", help="Export rich analysis to CSV/XLSX")
    export_parser.add_argument("--property-folder", required=True, help="Property config folder")
    export_parser.add_argument("--start-date", help="Start date (YYYY-MM-DD)")
    export_parser.add_argument("--end-date", help="End date (YYYY-MM-DD)")
    export_parser.add_argument("--as-of", default=None, help="Pace cutoff date (YYYY-MM-DD)")
    export_parser.add_argument("--listings", default=None, help="Comma-separated listing names to include")
    export_parser.add_argument("--since-go-live", action="store_true", help="Use each listing's go-live date as start")
    export_parser.add_argument("--periods", choices=("monthly", "weekly"), default=None, help="Break down by period")
    export_parser.add_argument("--format", choices=("csv", "xlsx"), default="csv", dest="export_format", help="Output format")
    export_parser.add_argument("--output", "-o", required=True, help="Output file or directory path")
    export_parser.add_argument("--data-root", default=str(DEFAULT_DATA_ROOT), help="Data root directory")
    export_parser.add_argument("--bands", default=DEFAULT_BANDS, help="Booking window bands")

    parser.add_argument("--csv", help="Path to booking CSV export")
    _add_run_arguments(parser)
    return parser


def _add_run_arguments(parser: argparse.ArgumentParser) -> None:
    parser.add_argument("--property-id", default="LAFAVE", help="Property ID label for output")
    parser.add_argument("--property-name", default="LaFave", help="Property name label for output")
    parser.add_argument(
        "--property-folder",
        default=None,
        help="Property config folder name (e.g. flohom, WMB); inferred from CSV parent dir if omitted",
    )
    parser.add_argument("--start-date", help="Start date (YYYY-MM-DD)")
    parser.add_argument("--end-date", help="End date (YYYY-MM-DD)")
    parser.add_argument(
        "--date-basis",
        choices=("stay", "arrival", "reservation"),
        default="stay",
        help=(
            "Date filter: stay = bookings with occupied nights in range (default); "
            "arrival = check-in in range; reservation = booked in range"
        ),
    )
    parser.add_argument(
        "--inventory-listings",
        type=int,
        default=30,
        help="Total listing count for portfolio occupancy/revpar denominator",
    )
    parser.add_argument(
        "--breakdown-by",
        choices=("listing", "grouping"),
        default="listing",
        help="Second section breakdown grain",
    )
    parser.add_argument(
        "--bands",
        default=DEFAULT_BANDS,
        help=f"Booking window bands, e.g. '{DEFAULT_BANDS}'",
    )
    parser.add_argument(
        "--format",
        choices=("text", "json"),
        default="text",
        help="Output format",
    )
    parser.add_argument(
        "--data-root",
        default=str(DEFAULT_DATA_ROOT),
        help="Data root directory (default: data)",
    )
    parser.add_argument(
        "--yoy-compare-start-date",
        default=None,
        help="YoY comparison period start (YYYY-MM-DD); for live_listings comparable units",
    )
    parser.add_argument(
        "--yoy-compare-end-date",
        default=None,
        help="YoY comparison period end (YYYY-MM-DD)",
    )


def _run_sync(args: argparse.Namespace) -> int:
    result = sync_properties(
        data_root=args.data_root,
        property_folder=args.property_folder,
    )
    if args.format == "json":
        print(json.dumps(result, indent=2))
        return 1 if result["errors"] else 0

    for item in result["properties"]:
        print(
            f"Synced {item['folder']} ({item['tab']}): "
            f"{item['row_count']} rows -> {item['cache_path']}"
        )
    for error in result["errors"]:
        print(f"Error syncing {error['folder']}: {error['error']}", file=sys.stderr)
    return 1 if result["errors"] else 0


def _run_snapshot(args: argparse.Namespace) -> int:
    if not args.start_date or not args.end_date:
        raise SystemExit("--start-date and --end-date are required")
    if not args.csv and not args.property_folder:
        raise SystemExit("Provide --csv or --property-folder")

    result = run_snapshot(
        start_date=args.start_date,
        end_date=args.end_date,
        csv_path=args.csv,
        property_id=args.property_id,
        property_name=args.property_name,
        date_basis=args.date_basis,
        breakdown_by=args.breakdown_by,
        bands=args.bands,
        inventory_listings=args.inventory_listings,
        property_folder=args.property_folder,
        data_root=args.data_root,
        yoy_compare_start_date=args.yoy_compare_start_date,
        yoy_compare_end_date=args.yoy_compare_end_date,
    )

    if args.format == "json":
        print(json.dumps(result.to_dict(), indent=2))
    else:
        print(
            multi_snapshot_to_text(
                portfolio=result.portfolio,
                listings=result.breakdown,
                invalid_row_count=result.invalid_rows_skipped,
                breakdown_label=result.breakdown_by.capitalize(),
            )
        )
    return 0


def _run_export(args: argparse.Namespace) -> int:
    from historical_snapshot.config import load_property_config
    from historical_snapshot.core.bands import parse_bands
    from historical_snapshot.core.metrics import (
        CANCELLED_STATUSES,
        build_booking_details,
        build_calendar_details,
        build_night_details,
    )
    from historical_snapshot.core.aggregations import (
        aggregate_blended,
        aggregate_by_channel,
        aggregate_by_date,
        aggregate_by_day_type,
        aggregate_by_dow,
        aggregate_by_period,
        aggregate_adr_stats,
    )
    from historical_snapshot.io.property_reader import read_property_bookings, resolve_property_data_path
    from historical_snapshot.service import resolve_available_room_nights, resolve_inventory_units

    config = load_property_config(args.property_folder)
    if config is None:
        print(f"Unknown property folder: {args.property_folder}", file=sys.stderr)
        return 1

    path = resolve_property_data_path(config, data_root=args.data_root)
    records, issues = read_property_bookings(path, property_config=config, data_root=args.data_root)
    live_dates = config.listing_live_dates or {}
    listing_inventory = config.listing_inventory or {}
    bands = parse_bands(args.bands) if isinstance(args.bands, str) else args.bands
    as_of = date.fromisoformat(args.as_of) if args.as_of else None

    if args.listings:
        listing_filter = {l.strip() for l in args.listings.split(",")}
        records = [r for r in records if r.listing_name in listing_filter]
        filtered_live_dates = {k: v for k, v in live_dates.items() if k in listing_filter}
    else:
        listing_filter = None
        filtered_live_dates = live_dates

    if args.since_go_live:
        if not filtered_live_dates:
            print("No listing live dates in config for --since-go-live", file=sys.stderr)
            return 1
        start = min(filtered_live_dates.values())
        end = date.today()
    else:
        if not args.start_date or not args.end_date:
            print("--start-date and --end-date required (or use --since-go-live)", file=sys.stderr)
            return 1
        start = date.fromisoformat(args.start_date)
        end = date.fromisoformat(args.end_date)

    non_cancelled = [r for r in records if r.status not in CANCELLED_STATUSES]

    nd = build_night_details(non_cancelled, start, end, date_basis="stay", as_of_date=as_of)
    bd = build_booking_details(non_cancelled, start, end, date_basis="stay", as_of_date=as_of)
    cd = build_calendar_details(nd, start, end, listing_live_dates=filtered_live_dates)

    arn = resolve_available_room_nights(
        start, end,
        inventory_mode=config.inventory_mode,
        inventory_units=None,
        listing_live_dates=filtered_live_dates or None,
        listing_inventory=listing_inventory or None,
    )
    inv = resolve_inventory_units(
        non_cancelled, start, end,
        inventory_mode=config.inventory_mode,
        manual_units=config.default_inventory_listings,
        listing_live_dates=filtered_live_dates or None,
        listing_inventory=listing_inventory or None,
    )

    blended = aggregate_blended(
        bd, nd, cd,
        property_id=config.property_id,
        property_name=config.property_name,
        start_date=start, end_date=end,
        bands=bands,
        available_room_nights=arn,
        inventory_units=inv,
        as_of_date=as_of,
    )
    dt_metrics = aggregate_by_day_type(nd, bd, cd)
    ch_metrics = aggregate_by_channel(bd, nd, cd)
    dow_metrics = aggregate_by_dow(nd, cd)
    date_metrics = aggregate_by_date(nd, cd)
    period_metrics = aggregate_by_period(nd, bd, cd, period=args.periods or "monthly", start_date=start, end_date=end) if args.periods else []

    def q(v):
        if v is None:
            return ""
        if isinstance(v, Decimal):
            return float(v.quantize(Decimal("0.01")))
        return v

    output_path = Path(args.output)

    if args.export_format == "xlsx":
        try:
            import openpyxl
        except ImportError:
            print("openpyxl required for xlsx export: pip install openpyxl", file=sys.stderr)
            return 1
        _export_xlsx(output_path, config, start, end, as_of, blended, dt_metrics, ch_metrics, dow_metrics, date_metrics, period_metrics, bd, nd, cd, q)
    else:
        output_path.mkdir(parents=True, exist_ok=True)
        _export_csvs(output_path, config, start, end, as_of, blended, dt_metrics, ch_metrics, dow_metrics, date_metrics, period_metrics, bd, nd, cd, q)

    return 0


def _write_csv(path: Path, fieldnames: list[str], rows: list[dict]) -> None:
    with path.open("w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=fieldnames)
        w.writeheader()
        w.writerows(rows)


def _export_csvs(out_dir, config, start, end, as_of, blended, dt_metrics, ch_metrics, dow_metrics, date_metrics, period_metrics, bd, nd, cd, q):
    _write_csv(out_dir / "01_summary.csv", [
        "property", "start_date", "end_date", "as_of", "bookings", "room_nights",
        "revenue", "adr_weighted", "adr_median", "adr_max", "occupancy_pct", "revpar",
    ], [{
        "property": config.property_name, "start_date": start, "end_date": end,
        "as_of": as_of or "", "bookings": blended.bookings_count,
        "room_nights": blended.room_nights_sold, "revenue": q(blended.room_revenue),
        "adr_weighted": q(blended.adr), "adr_median": q(blended.adr_median),
        "adr_max": q(blended.adr_max), "occupancy_pct": q(blended.occupancy_pct),
        "revpar": q(blended.revpar),
    }])

    dt_rows = []
    for dt_name in ("Sun-Thu", "Fri-Sat"):
        m = dt_metrics.get(dt_name)
        if m:
            dt_rows.append({
                "day_type": dt_name, "available_rn": m.available_room_nights,
                "bookings_touching": m.bookings_touching, "room_nights_sold": m.room_nights_sold,
                "revenue": q(m.room_revenue), "adr_weighted": q(m.adr_weighted),
                "adr_median": q(m.adr_stats.median), "adr_mean": q(m.adr_stats.mean),
                "adr_max": q(m.adr_stats.max), "occupancy_pct": q(m.occupancy_pct),
                "revpar": q(m.revpar),
            })
    dt_rows.append({
        "day_type": "Blended", "available_rn": sum(1 for c in cd if c.is_available),
        "bookings_touching": blended.bookings_count, "room_nights_sold": blended.room_nights_sold,
        "revenue": q(blended.room_revenue), "adr_weighted": q(blended.adr),
        "adr_median": q(blended.adr_median), "adr_mean": q(blended.adr),
        "adr_max": q(blended.adr_max), "occupancy_pct": q(blended.occupancy_pct),
        "revpar": q(blended.revpar),
    })
    _write_csv(out_dir / "02_weekday_weekend.csv", [
        "day_type", "available_rn", "bookings_touching", "room_nights_sold",
        "revenue", "adr_weighted", "adr_median", "adr_mean", "adr_max", "occupancy_pct", "revpar",
    ], dt_rows)

    ch_rows = []
    for ch, m in ch_metrics.items():
        ch_rows.append({
            "channel": ch, "bookings": m.bookings_count, "room_nights": m.room_nights,
            "revenue": q(m.room_revenue), "adr_weighted": q(m.adr_weighted),
            "adr_median": q(m.adr_stats.median), "adr_mean": q(m.adr_stats.mean),
            "adr_max": q(m.adr_stats.max), "revenue_share_pct": q(m.revenue_share_pct),
            "nights_share_pct": q(m.nights_share_pct),
            "revpar_of_total_available": q(m.revpar_of_total_available),
        })
    _write_csv(out_dir / "03_channel_mix.csv", [
        "channel", "bookings", "room_nights", "revenue", "adr_weighted",
        "adr_median", "adr_mean", "adr_max", "revenue_share_pct",
        "nights_share_pct", "revpar_of_total_available",
    ], ch_rows)

    dow_rows = []
    from historical_snapshot.core.metrics import DOW_LABELS
    for dow in DOW_LABELS:
        m = dow_metrics.get(dow)
        if m:
            dow_rows.append({
                "day_of_week": dow, "available_rn": m.available_room_nights,
                "room_nights_sold": m.room_nights_sold, "revenue": q(m.room_revenue),
                "adr_weighted": q(m.adr_weighted), "occupancy_pct": q(m.occupancy_pct),
                "revpar": q(m.revpar),
            })
    _write_csv(out_dir / "04_dow_grid.csv", [
        "day_of_week", "available_rn", "room_nights_sold", "revenue",
        "adr_weighted", "occupancy_pct", "revpar",
    ], dow_rows)

    if period_metrics:
        p_rows = [{
            "period": p.period_label, "start_date": p.start_date, "end_date": p.end_date,
            "available_rn": p.available_room_nights, "bookings": p.bookings_count,
            "room_nights_sold": p.room_nights_sold, "revenue": q(p.room_revenue),
            "adr_weighted": q(p.adr_weighted), "adr_median": q(p.adr_stats.median),
            "occupancy_pct": q(p.occupancy_pct), "revpar": q(p.revpar),
        } for p in period_metrics]
        _write_csv(out_dir / "05_periods.csv", [
            "period", "start_date", "end_date", "available_rn", "bookings",
            "room_nights_sold", "revenue", "adr_weighted", "adr_median",
            "occupancy_pct", "revpar",
        ], p_rows)

    night_rows = [{
        "night_date": n.night_date, "day_of_week": n.day_of_week, "day_type": n.day_type,
        "listing": n.listing_name, "channel": n.channel,
        "prorated_revenue": q(n.prorated_revenue),
        "check_in": n.check_in_date, "check_out": n.check_out_date,
    } for n in nd]
    _write_csv(out_dir / "06_night_level.csv", [
        "night_date", "day_of_week", "day_type", "listing", "channel",
        "prorated_revenue", "check_in", "check_out",
    ], night_rows)

    booking_rows = [{
        "listing": b.listing_name, "channel": b.channel,
        "reservation_date": b.reservation_date, "check_in": b.check_in_date,
        "check_out": b.check_out_date, "nights_in_range": b.nights_in_range,
        "revenue_in_range": q(b.revenue_in_range), "booking_adr": q(b.booking_adr),
        "booking_window_days": b.booking_window_days,
        "day_types_touched": ",".join(sorted(b.day_types_touched)),
    } for b in bd]
    _write_csv(out_dir / "07_booking_detail.csv", [
        "listing", "channel", "reservation_date", "check_in", "check_out",
        "nights_in_range", "revenue_in_range", "booking_adr",
        "booking_window_days", "day_types_touched",
    ], booking_rows)

    cal_rows = [{
        "date": c.night_date, "day_of_week": c.day_of_week, "day_type": c.day_type,
        "listing": c.listing_name, "is_available": int(c.is_available),
        "is_sold": int(c.is_sold), "revenue": q(c.revenue),
    } for c in cd]
    _write_csv(out_dir / "08_calendar.csv", [
        "date", "day_of_week", "day_type", "listing", "is_available", "is_sold", "revenue",
    ], cal_rows)

    date_rows = [{
        "date": d.night_date, "day_of_week": d.day_of_week, "day_type": d.day_type,
        "available_rn": d.available_room_nights, "room_nights_sold": d.room_nights_sold,
        "revenue": q(d.room_revenue), "occupancy_pct": q(d.occupancy_pct),
    } for d in date_metrics]
    _write_csv(out_dir / "09_date_heatmap.csv", [
        "date", "day_of_week", "day_type", "available_rn", "room_nights_sold",
        "revenue", "occupancy_pct",
    ], date_rows)

    print(f"Exported {len(list(out_dir.glob('*.csv')))} CSVs to {out_dir}")


def _export_xlsx(out_path, config, start, end, as_of, blended, dt_metrics, ch_metrics, dow_metrics, date_metrics, period_metrics, bd, nd, cd, q):
    import openpyxl
    wb = openpyxl.Workbook()

    def _add_sheet(name, headers, rows):
        ws = wb.create_sheet(name)
        ws.append(headers)
        for row in rows:
            ws.append([row.get(h, "") for h in headers])

    wb.remove(wb.active)

    _add_sheet("Summary", [
        "property", "start_date", "end_date", "as_of", "bookings", "room_nights",
        "revenue", "adr_weighted", "adr_median", "adr_max", "occupancy_pct", "revpar",
    ], [{
        "property": config.property_name, "start_date": str(start), "end_date": str(end),
        "as_of": str(as_of) if as_of else "", "bookings": blended.bookings_count,
        "room_nights": blended.room_nights_sold, "revenue": q(blended.room_revenue),
        "adr_weighted": q(blended.adr), "adr_median": q(blended.adr_median),
        "adr_max": q(blended.adr_max), "occupancy_pct": q(blended.occupancy_pct),
        "revpar": q(blended.revpar),
    }])

    dt_rows = []
    for dt_name in ("Sun-Thu", "Fri-Sat"):
        m = dt_metrics.get(dt_name)
        if m:
            dt_rows.append({
                "day_type": dt_name, "available_rn": m.available_room_nights,
                "bookings_touching": m.bookings_touching, "room_nights_sold": m.room_nights_sold,
                "revenue": q(m.room_revenue), "adr_weighted": q(m.adr_weighted),
                "adr_median": q(m.adr_stats.median), "adr_max": q(m.adr_stats.max),
                "occupancy_pct": q(m.occupancy_pct), "revpar": q(m.revpar),
            })
    _add_sheet("Weekday_Weekend", [
        "day_type", "available_rn", "bookings_touching", "room_nights_sold",
        "revenue", "adr_weighted", "adr_median", "adr_max", "occupancy_pct", "revpar",
    ], dt_rows)

    ch_rows = [{
        "channel": ch, "bookings": m.bookings_count, "room_nights": m.room_nights,
        "revenue": q(m.room_revenue), "adr_weighted": q(m.adr_weighted),
        "adr_median": q(m.adr_stats.median), "adr_max": q(m.adr_stats.max),
        "revenue_share_pct": q(m.revenue_share_pct),
    } for ch, m in ch_metrics.items()]
    _add_sheet("Channel_Mix", [
        "channel", "bookings", "room_nights", "revenue", "adr_weighted",
        "adr_median", "adr_max", "revenue_share_pct",
    ], ch_rows)

    if period_metrics:
        _add_sheet("Periods", [
            "period", "start_date", "end_date", "available_rn", "bookings",
            "room_nights_sold", "revenue", "adr_weighted", "adr_median",
            "occupancy_pct", "revpar",
        ], [{
            "period": p.period_label, "start_date": str(p.start_date), "end_date": str(p.end_date),
            "available_rn": p.available_room_nights, "bookings": p.bookings_count,
            "room_nights_sold": p.room_nights_sold, "revenue": q(p.room_revenue),
            "adr_weighted": q(p.adr_weighted), "adr_median": q(p.adr_stats.median),
            "occupancy_pct": q(p.occupancy_pct), "revpar": q(p.revpar),
        } for p in period_metrics])

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_path))
    print(f"Exported to {out_path}")


def main() -> int:
    parser = build_parser()
    args = parser.parse_args()
    if args.command == "sync":
        return _run_sync(args)
    if args.command == "run":
        return _run_snapshot(args)
    if args.command == "export":
        return _run_export(args)
    if args.start_date and args.end_date and (args.csv or args.property_folder):
        return _run_snapshot(args)
    parser.print_help()
    return 2


if __name__ == "__main__":
    raise SystemExit(main())
